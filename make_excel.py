import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Weekly Progress"
ws.views.sheetView[0].showGridLines = True

# Palettes (Google Sheets style)
PRIMARY_COLOR = "1A73E8"
HEADER_BG = "E8F0FE"
HEADER_FG = "1967D2"
ZEBRA_BG = "F8F9FA"
BORDER_COLOR = "DADCE0"

thin_side = Side(style='thin', color=BORDER_COLOR)
cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

# Title & Metadata
ws['A1'] = "WCTE Project Weekly Progress Tracker"
ws['A1'].font = Font(name="Arial", size=14, bold=True, color="137333")

ws['A2'] = "Arkaprava Ghosh - DA25C003"
ws['A2'].font = Font(name="Arial", size=11, bold=True, color="202124")

ws['A3'] = "Reporting to: Ramesh sir (CTO) and Kavitha mam (Tech Lead)"
ws['A3'].font = Font(name="Arial", size=10, italic=True, color="5F6368")

ws['A4'] = "Reporting period: 24 August–25 October 2026 | Updated: 21 October 2026"
ws['A4'].font = Font(name="Arial", size=10, italic=True, color="5F6368")

# Header row at Row 6
headers = ["Week and date range", "Work completed / in progress", "Bottlenecks faced"]
for col_idx, header in enumerate(headers, start=1):
    cell = ws.cell(row=6, column=col_idx, value=header)
    cell.font = Font(name="Arial", size=11, bold=True, color=HEADER_FG)
    cell.fill = PatternFill(start_color=HEADER_BG, end_color=HEADER_BG, fill_type="solid")
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    cell.border = cell_border
ws.row_dimensions[6].height = 26

# Data rows
data = [
    (
        "Week 1\n24–30 August 2026",
        "• Met Ramesh sir and Kavitha mam on 27 August.\n• Discussed practical ML problems, factory automation challenges, and possible project directions.",
        "• Needed to finalise a suitable, high-impact problem statement."
    ),
    (
        "Week 2\n31 August–6 September 2026",
        "• Reviewed the existing VisureX approach and project poster.\n• Researched operator logbook digitization for MSMEs.\n• Drafted and submitted initial proposal on 2 September: Hands-free logbook & production tracking using Meta Glasses, ArUco marker gaze locks, mobile edge proxy, and CV dial/tray reading.",
        "• Changes in ambient factory lighting and camera orientation affected reading stability.\n• Wearable hardware constraints and edge processing reliability raised feasibility questions."
    ),
    (
        "Week 3\n7–13 September 2026",
        "• Explored Meta glasses workflow feasibility vs. fixed industrial inspection.\n• Researched optical capture bounds, Bluetooth bandwidth limits, and display parsing models.\n• Coordinated technical review on MSME shop-floor ergonomics and deployment costs.",
        "• Capture latency and battery constraints in harsh industrial settings.\n• Need to evaluate whether wearable glasses or fixed-surface visual inspection provides higher industry demand."
    ),
    (
        "Week 4\n14–20 September 2026",
        "• Pivoted project scope toward automated flat surface quality inspection.\n• Submitted revised project abstract on 15 September: Integrated visual inspection platform combining 2D surface anomaly detection with Automated Optical Inspection (AOI) for solar panels and sheet materials.\n• Defined the dual-engine concept: classical CV for preprocessing and ML for defect segmentation.",
        "• Narrowing down specific anomaly classes and defining quantifiable severity criteria.\n• Limited public availability of pixel-level defect masks."
    ),
    (
        "Week 5\n21–27 September 2026",
        "• Conducted literature review on solar photovoltaic (PV) soiling and defect inspection.\n• Formulated project boundary: in-scope stains (dust/soiling, bird droppings, mud) vs. out-of-scope defects (snow, mechanical cracks, electrical burns).\n• Prepared initial architectural proposal comparing single-stage classifiers against two-stage segmentation.",
        "• Scheduled review meeting was postponed due to team travel.\n• Most public datasets only provide image-level labels rather than pixel-level segmentation masks."
    ),
    (
        "Week 6\n28 September–4 October 2026",
        "• Attended technical review on 28 September with Ramesh sir and Kavitha mam.\n• Finalized the two-stage inspection architecture: Stage A for panel localization/rectification and Stage B for stain segmentation.\n• Defined explainability goals: replacing black-box classification with exact stain coverage % and localized hotspot alerts.",
        "• Background artifacts (sky, roof tiles, shadows) distort coverage metrics if panels are not strictly rectified.\n• Need for a robust boundary-detection baseline."
    ),
    (
        "Week 7\n5–11 October 2026",
        "• Attended meeting at WCTE Centre on 7 October discussing plausible datasets for model training.\n• Evaluated public repositories: Kaggle (Faulty Solar Panel Images), Mendeley Data (Solar Panel Soiling), and Roboflow Universe.\n• Designed data-hygiene protocol: Perceptual Hashing (pHash) to detect duplicate web images and prevent train–test leakage.",
        "• Reused images across public repositories cause severe train–test data leakage.\n• Extreme class imbalance between large clean areas and small, high-risk bird droppings."
    ),
    (
        "Week 8\n12–18 October 2026",
        "• Set up local development environment (Python 3.12, PyTorch, Albumentations, SMP, OpenCV).\n• Finalized operational decision rules: High Stain if Coverage ≥ 10% or any single bird blob ≥ 2%; 4×4 sub-panel grid breakdown.\n• Implemented Stage A perspective homography rectifier to canonical 512×512 view.\n• Formulated dataset curation pipeline: pHash deduplication, 70/15/15 group-stratified splits, and SAM 2/CVAT-aligned ground-truth mask preparation.",
        "• Ensuring augmentations avoid hue shifts that misrepresent dust tones.\n• Adapting models and training pipelines for efficient CPU execution."
    ),
    (
        "Week 9 — ongoing\n19–25 October 2026",
        "• Ingested and standardized multi-source datasets (Kaggle, Mendeley, Roboflow) into unified manifest.\n• Model is currently in the training process for the computer vision solar panel stain detection and segmentation network.\n• Monitoring training iterations, batch scheduling, and convergence across initial epochs.",
        "• Training execution time on local workstation hardware.\n• Balancing batch size and memory footprint during segmentation training."
    )
]

start_row = 7
for idx, (col1, col2, col3) in enumerate(data):
    current_row = start_row + idx
    is_even = (idx % 2 == 1)
    row_fill = PatternFill(start_color=ZEBRA_BG, end_color=ZEBRA_BG, fill_type="solid") if is_even else None

    # Col A: Week
    c1 = ws.cell(row=current_row, column=1, value=col1)
    c1.font = Font(name="Arial", size=10, bold=True, color="202124")
    c1.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    c1.border = cell_border
    if row_fill: c1.fill = row_fill

    # Col B: Work Done
    c2 = ws.cell(row=current_row, column=2, value=col2)
    c2.font = Font(name="Arial", size=10, color="202124")
    c2.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    c2.border = cell_border
    if row_fill: c2.fill = row_fill

    # Col C: Bottlenecks
    c3 = ws.cell(row=current_row, column=3, value=col3)
    c3.font = Font(name="Arial", size=10, color="3C4043")
    c3.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    c3.border = cell_border
    if row_fill: c3.fill = row_fill

    # Row height based on text lines
    max_lines = max(col1.count('\n') + 1, col2.count('\n') + 1, col3.count('\n') + 1)
    ws.row_dimensions[current_row].height = max(45, max_lines * 18)

# Set Column Widths
ws.column_dimensions['A'].width = 24
ws.column_dimensions['B'].width = 68
ws.column_dimensions['C'].width = 46

out_path = r"C:\Users\ghosh\.gemini\antigravity\scratch\solar_panel_stain_detection\WCTE_Weekly_Progress_Tracker_Arkaprava_Ghosh.xlsx"
wb.save(out_path)
print(f"Saved: {out_path}")
